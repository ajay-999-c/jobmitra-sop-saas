"""Conservative, insert-only import of the Phase 1A workbook. No schema operations."""

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path

from sqlalchemy import select, text
from sqlalchemy.orm import Session

from jobmitra.cli import database_failure_details
from jobmitra.core.config import get_settings
from jobmitra.core.database import build_engine
from jobmitra.modules.companies.models import Company, CompanyContact
from jobmitra.modules.companies.normalization import company_name_key

HEADERS = {
    "Companies": (
        "Company name",
        "Company name key",
        "Source Register IDs",
        "General phone numbers",
        "Address line",
        "City",
        "State",
        "Postal code",
        "Review status",
        "Source Excel rows",
        "Original name spellings",
    ),
    "Contacts": (
        "Company name key",
        "Contact name / source text",
        "Designation",
        "Phone numbers",
        "Source Register IDs",
        "Source Excel rows",
        "Review status",
    ),
    "Review": (
        "Review type",
        "Company name keys",
        "Source Excel rows",
        "Reason",
        "Decision",
        "Reviewer note",
    ),
    "Source Rows": (
        "Original Excel row",
        "Register ID",
        "Company / client (as read)",
        "Contact person(s) (as read)",
        "Phone 1 - readable digits",
        "Phone 2 - readable digits",
    ),
}
UNCLEAR = re.compile(r"\b(?:unclear|unreadable|illegible|unknown)\b|\?", re.I)
SPLIT = re.compile(r"[/&;,]|\band\b", re.I)
ROLE_ONLY = re.compile(r"^(?:mr|mrs|ms|dr|hr|manager|owner|director|proprietor|sir)[.\s]*$", re.I)
IMPORT_LOCK = 48650168417482


class InvalidWorkbook(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise InvalidWorkbook(message)


def array(value, location):
    try:
        result = json.loads(value)
    except (ValueError, TypeError):
        raise InvalidWorkbook(f"{location}: expected a JSON string array") from None
    require(isinstance(result, list), f"{location}: expected an array")
    require(
        all(isinstance(v, str) and v.strip() and v == v.strip() for v in result),
        f"{location}: array elements must be nonblank strings without outer whitespace",
    )
    require(len(set(result)) == len(result), f"{location}: repeated array elements")
    return result


def row_numbers(value, location):
    require(isinstance(value, str), f"{location}: expected source row references")
    parts = [p.strip() for p in value.split(";")]
    require(
        all(p.isdecimal() and int(p) >= 2 for p in parts), f"{location}: invalid row references"
    )
    result = [int(p) for p in parts]
    require(len(set(result)) == len(result), f"{location}: duplicate row references")
    return result


def optional_text(value, location):
    require(value is None or isinstance(value, str), f"{location}: expected text or blank")
    return value.strip() or None if value is not None else None


def read_inputs(source: Path, original: Path):
    import openpyxl  # Optional import extra; the backend does not need an Excel reader.

    snapshots = {"cleaned": source.read_bytes(), "original": original.read_bytes()}
    tables = {}
    with_workbook = openpyxl.load_workbook(
        BytesIO(snapshots["cleaned"]), read_only=True, data_only=False, keep_links=False
    )
    try:
        require(set(with_workbook.sheetnames) == set(HEADERS), "Unexpected cleaned workbook sheets")
        for sheet_name, headers in HEADERS.items():
            rows = list(with_workbook[sheet_name].iter_rows())
            require(len(rows) >= 6, f"{sheet_name}: missing table header")
            require(tuple(c.value for c in rows[5]) == headers, f"{sheet_name}: unexpected columns")
            require(
                not any(c.data_type in {"f", "e"} for row in rows for c in row),
                f"{sheet_name}: formulas and error cells are not allowed",
            )
            tables[sheet_name] = [
                {**dict(zip(headers, (c.value for c in row), strict=True)), "_row": number}
                for number, row in enumerate(rows[6:], 7)
                if any(c.value is not None for c in row)
            ]
    finally:
        with_workbook.close()
    workbook = openpyxl.load_workbook(
        BytesIO(snapshots["original"]), read_only=True, data_only=False, keep_links=False
    )
    try:
        require(workbook.sheetnames == ["Sheet1"], "Unexpected original workbook sheets")
        rows = list(workbook["Sheet1"].iter_rows())
        require(
            tuple(c.value for c in rows[0]) == HEADERS["Source Rows"][1:],
            "Unexpected original workbook columns",
        )
        require(
            not any(c.data_type in {"f", "e"} for row in rows for c in row),
            "Original workbook contains formulas or error cells",
        )
        originals = {
            i: tuple(c.value for c in row)
            for i, row in enumerate(rows[1:], 2)
            if any(c.value is not None for c in row)
        }
    finally:
        workbook.close()
    hashes = {key: hashlib.sha256(value).hexdigest() for key, value in snapshots.items()}
    return tables, originals, hashes


@dataclass
class ContactCandidate:
    name: str
    designation: str | None
    row: int
    reasons: list[str] = field(default_factory=list)


@dataclass
class CompanyCandidate:
    name: str
    key: str
    source_ids: list[str]
    phones: list[str]
    address: dict[str, str | None]
    row: int
    contacts: list[ContactCandidate] = field(default_factory=list)
    reasons: set[str] = field(default_factory=set)


@dataclass
class ImportPlan:
    companies: list[CompanyCandidate]
    source_entries: int
    review_items: int
    hashes: dict[str, str] = field(default_factory=dict)

    def summary(self):
        accepted = [c for c in self.companies if not c.reasons]
        contacts = [contact for c in accepted for contact in c.contacts if not contact.reasons]
        return {
            "company_candidates": len(self.companies),
            "original_entries": self.source_entries,
            "original_register_ids": sum(len(c.source_ids) for c in self.companies),
            "eligible_companies": len(accepted),
            "withheld_companies": len(self.companies) - len(accepted),
            "eligible_source_register_ids": sum(len(c.source_ids) for c in accepted),
            "eligible_general_phone_values": sum(len(c.phones) for c in accepted),
            "eligible_contacts": len(contacts),
            "contacts_skipped_ambiguous": sum(
                bool(p.reasons) for c in accepted for p in c.contacts
            ),
            "contacts_skipped_parent_withheld": sum(
                len(c.contacts) for c in self.companies if c.reasons
            ),
            "contact_candidates": sum(len(c.contacts) for c in self.companies),
            "review_items": self.review_items,
        }


def validate_tables(tables, originals, *, expected_companies=166, expected_entries=180):
    require(len(tables["Companies"]) == expected_companies, "Unexpected company candidate count")
    require(
        len(tables["Source Rows"]) == len(originals) == expected_entries,
        "Unexpected original entry count",
    )
    sources = {}
    for row in tables["Source Rows"]:
        number = row["Original Excel row"]
        values = tuple(row[h] for h in HEADERS["Source Rows"][1:])
        require(
            type(number) is int and number in originals and originals[number] == values,
            f"Source Rows!{row['_row']}: does not match original workbook",
        )
        register, name, contact, *phones = values
        require(
            isinstance(register, str) and register.strip() == register and bool(register),
            "Invalid original Register ID",
        )
        require(register not in sources, "Repeated original Register ID")
        require(isinstance(name, str) and bool(name.strip()), "Missing original company name")
        require(contact is None or isinstance(contact, str), "Invalid original contact type")
        require(
            all(p is None or (isinstance(p, str) and p.isdecimal()) for p in phones),
            "Original phone values must remain digit strings",
        )
        sources[register] = {
            "number": number,
            "key": company_name_key(name),
            "name": name,
            "contact": contact,
            "phones": [p for p in phones if p],
        }
    require(
        {r["number"] for r in sources.values()} == set(originals),
        "Incomplete original row coverage",
    )
    companies, assigned_ids = {}, []
    for row in tables["Companies"]:
        location = f"Companies!{row['_row']}"
        require(
            isinstance(row["Company name"], str) and bool(row["Company name"].strip()),
            f"{location}: missing company name",
        )
        key = company_name_key(row["Company name"])
        require(
            key == row["Company name key"] and key not in companies,
            f"{location}: inconsistent or duplicate normalized name key",
        )
        ids = array(row["Source Register IDs"], location)
        require(bool(ids) and all(rid in sources for rid in ids), f"{location}: unknown source IDs")
        group = [sources[rid] for rid in ids]
        require(all(r["key"] == key for r in group), f"{location}: incorrect source ID mapping")
        require(
            row["Company name"] in {r["name"] for r in group},
            f"{location}: display name is not an original spelling",
        )
        require(
            set(row_numbers(row["Source Excel rows"], location)) == {r["number"] for r in group},
            f"{location}: incorrect source row mapping",
        )
        phones = array(row["General phone numbers"], location)
        require(
            set(phones) == {p for r in group for p in r["phones"]},
            f"{location}: distinct source phone information was lost or added",
        )
        address = {
            name: optional_text(row[column], location)
            for name, column in (
                ("address_line", "Address line"),
                ("city", "City"),
                ("state", "State"),
                ("postal_code", "Postal code"),
            )
        }
        companies[key] = CompanyCandidate(
            row["Company name"], key, ids, phones, address, row["_row"]
        )
        if UNCLEAR.search(row["Company name"]):
            companies[key].reasons.add("unclear_company_name")
        assigned_ids.extend(ids)
    require(
        Counter(assigned_ids) == Counter(sources.keys()),
        "Source IDs must occur exactly once across companies",
    )

    reviews = {}
    for row in tables["Review"]:
        location = f"Review!{row['_row']}"
        kind = row["Review type"]
        require(
            kind in {"Merge review", "Similar names", "Shared phone", "Contact split"},
            f"{location}: unknown review type",
        )
        require(isinstance(row["Company name keys"], str), f"{location}: missing company keys")
        keys = frozenset(row["Company name keys"].split("; "))
        require(bool(keys) and keys <= companies.keys(), f"{location}: unknown company keys")
        refs = frozenset(row_numbers(row["Source Excel rows"], location))
        require(
            refs <= {r["number"] for r in sources.values() if r["key"] in keys},
            f"{location}: source rows do not belong to named companies",
        )
        identity = (kind, keys, refs)
        require(identity not in reviews, f"{location}: repeated review entry")
        decisions = {
            "Merge review": {"Merge confirmed"},
            "Similar names": {"Keep separate", "Verified source"},
            "Shared phone": {"Keep separate", "Verified source"},
        }
        approved = row["Decision"] in decisions.get(kind, set()) and bool(
            optional_text(row["Reviewer note"], location)
        )
        reviews[identity] = approved

    def identity_risk(kind, keys, group):
        refs = frozenset(r["number"] for r in group)
        if not reviews.get((kind, frozenset(keys), refs), False):
            for key in keys:
                companies[key].reasons.add("unresolved_" + kind.lower().replace(" ", "_"))

    # Recompute risks from the original entries; deleting flags cannot silently approve a merge.
    grouped = {key: [sources[rid] for rid in c.source_ids] for key, c in companies.items()}
    for key, group in grouped.items():
        payloads = {
            (r["key"], company_name_key(r["contact"]) if r["contact"] else None, tuple(r["phones"]))
            for r in group
        }
        if len(payloads) > 1:
            identity_risk("Merge review", [key], group)
    keys = list(companies)
    for i, key in enumerate(keys):
        for other in keys[i + 1 :]:
            names = [re.sub(r"[^\w]", "", k) for k in (key, other)]
            if SequenceMatcher(None, *names).ratio() >= 0.82:
                identity_risk("Similar names", [key, other], grouped[key] + grouped[other])
    phone_groups = defaultdict(list)
    for entry in sources.values():
        for phone in entry["phones"]:
            phone_groups[phone].append(entry)
    for group in phone_groups.values():
        keys = {entry["key"] for entry in group}
        if len(keys) > 1:
            identity_risk("Shared phone", keys, group)

    actual_contacts = set()
    for row in tables["Contacts"]:
        location = f"Contacts!{row['_row']}"
        key, name = row["Company name key"], row["Contact name / source text"]
        require(
            key in companies and isinstance(name, str) and bool(name.strip()),
            f"{location}: missing company or contact text",
        )
        require((key, name) not in actual_contacts, f"{location}: repeated contact record")
        actual_contacts.add((key, name))
        ids = array(row["Source Register IDs"], location)
        expected = [
            entry for entry in grouped[key] if entry["contact"] and entry["contact"].strip() == name
        ]
        require(
            bool(expected)
            and set(ids) == {rid for rid in companies[key].source_ids if sources[rid] in expected},
            f"{location}: incorrect contact source mapping",
        )
        require(
            set(row_numbers(row["Source Excel rows"], location)) == {r["number"] for r in expected},
            f"{location}: incorrect contact row mapping",
        )
        require(
            array(row["Phone numbers"], location) == [],
            f"{location}: source establishes no contact phone ownership",
        )
        reasons = []
        if SPLIT.search(name) or "REVIEW" in str(row["Review status"]).upper():
            reasons.append("ambiguous_contact_split")
        if UNCLEAR.search(name) or ROLE_ONLY.fullmatch(name.strip()):
            reasons.append("unclear_contact_name")
        companies[key].contacts.append(
            ContactCandidate(
                name, optional_text(row["Designation"], location), row["_row"], reasons
            )
        )
    expected_contacts = {(r["key"], r["contact"].strip()) for r in sources.values() if r["contact"]}
    require(
        actual_contacts == expected_contacts, "Contact text coverage differs from original entries"
    )
    return ImportPlan(list(companies.values()), len(sources), len(reviews))


def load_plan(source: Path, original: Path) -> ImportPlan:
    tables, originals, hashes = read_inputs(source, original)
    plan = validate_tables(tables, originals)
    plan.hashes = hashes
    return plan


def execute_import(session: Session, plan: ImportPlan, *, apply: bool) -> dict:
    """Caller owns the atomic transaction. Existing ORM entities are never attached or edited."""
    if apply:
        session.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": IMPORT_LOCK})
    existing = list(
        session.execute(select(Company.id, Company.company_name_key, Company.source_register_ids))
    )
    by_name = {r.company_name_key: r.id for r in existing}
    owners = defaultdict(set)
    existing_ids = {}
    for row in existing:
        existing_ids[row.id] = set(row.source_register_ids)
        for rid in row.source_register_ids:
            owners[rid].add(row.id)
    actions, pending = [], []
    for candidate in plan.companies:
        action = {
            "company_name": candidate.name,
            "company_name_key": candidate.key,
            "company_row": candidate.row,
            "source_register_ids": candidate.source_ids,
        }
        source_owners = set().union(*(owners[rid] for rid in candidate.source_ids))
        if candidate.reasons:
            action.update(status="withheld_ambiguous_company", reasons=sorted(candidate.reasons))
        elif source_owners:
            complete = (
                len(source_owners) == 1
                and set(candidate.source_ids) <= existing_ids[next(iter(source_owners))]
            )
            action.update(status="already_present" if complete else "existing_source_conflict")
        elif candidate.key in by_name:
            action.update(status="existing_name_conflict")
        else:
            contacts = [contact for contact in candidate.contacts if not contact.reasons]
            action.update(status="insert" if apply else "would_insert", contact_count=len(contacts))
            if apply:
                company = Company(
                    company_name=candidate.name,
                    general_phone_numbers=list(candidate.phones),
                    source_register_ids=list(candidate.source_ids),
                    **candidate.address,
                    contacts=[
                        CompanyContact(
                            contact_name=p.name, designation=p.designation, phone_numbers=[]
                        )
                        for p in contacts
                    ],
                )
                session.add(company)
                pending.append((company, action))
        actions.append(action)
    if apply:
        session.flush()
        for company, action in pending:
            action["company_id"] = str(company.id)
    totals = Counter(action["status"] for action in actions)
    insertions = [a for a in actions if a["status"] in {"insert", "would_insert"}]
    return {
        "mode": "apply" if apply else "dry-run",
        "validation": plan.summary(),
        "database_plan": dict(totals),
        "companies_to_insert": len(insertions),
        "contacts_to_insert": sum(a["contact_count"] for a in insertions),
        "inserted_companies": len(pending),
        "inserted_contacts": sum(a["contact_count"] for _, a in pending),
        "company_actions": actions,
        "contact_review": [
            {
                "company_row": c.row,
                "contact_row": p.row,
                "reasons": p.reasons or ["parent_company_withheld"],
            }
            for c in plan.companies
            for p in c.contacts
            if c.reasons or p.reasons
        ],
        "source_sha256": plan.hashes,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate or insert unambiguous Phase 1A company records"
    )
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--apply", action="store_true")
    parser.add_argument(
        "--source", type=Path, default=Path("outputs/phase-1a/company-dataset-cleaned.xlsx")
    )
    parser.add_argument("--original", type=Path, default=Path("company-dataset.xlsx"))
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = args.report or Path("outputs/phase-1c") / (
        datetime.now(UTC).strftime("%Y%m%dT%H%M%S%fZ")
        + ("-apply.json" if args.apply else "-dry-run.json")
    )
    engine = None
    try:
        require(not report.exists(), "Report already exists; choose a new path")
        report.parent.mkdir(parents=True, exist_ok=True)
        plan = load_plan(args.source, args.original)
        settings = get_settings()
        require(
            settings.db_name == "jobmitra_dev"
            and settings.db_host in {"localhost", "127.0.0.1", "::1"},
            "Importer requires local jobmitra_dev settings",
        )
        engine = build_engine(settings)
        with Session(engine) as session, session.begin():
            if args.dry_run:
                session.execute(text("SET TRANSACTION READ ONLY"))
            require(
                session.scalar(text("SELECT current_database()")) == "jobmitra_dev",
                "Unexpected database target",
            )
            require(
                session.scalar(text("SELECT current_schema()")) == "public", "Unexpected schema"
            )
            require(
                session.connection().dialect.server_version_info[0] == 18, "PostgreSQL 18 required"
            )
            require(
                session.execute(text("SELECT version_num FROM alembic_version")).scalar_one()
                == "0001_company_management",
                "Unexpected Alembic revision",
            )
            result = execute_import(session, plan, apply=args.apply)
    except InvalidWorkbook as error:
        print(json.dumps({"status": "validation_failed", "message": str(error)}))
        return 1
    except Exception as error:
        print(
            json.dumps(
                {
                    "status": "failed",
                    "message": "Import did not commit; transaction rolled back.",
                    **database_failure_details(error),
                }
            )
        )
        return 1
    finally:
        if engine is not None:
            engine.dispose()
    try:
        with report.open("x", encoding="utf-8") as output:
            output.write(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
        report.chmod(0o600)
    except OSError:
        print(
            json.dumps(
                {
                    "status": "committed_without_report" if args.apply else "report_failed",
                    "message": "Database work finished, but the report could not be written.",
                    **{
                        k: v
                        for k, v in result.items()
                        if k not in {"company_actions", "contact_review"}
                    },
                }
            )
        )
        return 2
    print(
        json.dumps(
            {
                "status": "ok",
                "report": str(report),
                **{
                    k: v
                    for k, v in result.items()
                    if k not in {"company_actions", "contact_review", "source_sha256"}
                },
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
